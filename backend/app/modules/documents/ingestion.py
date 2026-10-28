import io
import math
import re
from dataclasses import dataclass
from typing import Protocol
from zipfile import BadZipFile

from app.modules.documents.uploads import validate_document

MAX_EXTRACTED_CHARACTERS = 4_000_000
MAX_PDF_PAGES = 2_000
MAX_CHUNK_CHARACTERS = 1_200
CHUNK_OVERLAP_CHARACTERS = 160
MAX_CHUNKS_PER_DOCUMENT = 10_000


class DocumentProcessingError(Exception):
    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class ExtractedPage:
    number: int | None
    text: str


@dataclass(frozen=True)
class TextChunk:
    chunk_index: int
    content: str
    page_start: int | None
    page_end: int | None


class EmbeddingProvider(Protocol):
    model: str

    def embed(self, texts: list[str]) -> list[list[float]]: ...


def extract_document(filename: str, mime_type: str, data: bytes) -> list[ExtractedPage]:
    try:
        _, detected_mime = validate_document(filename, data)
    except Exception as error:
        if hasattr(error, "status_code"):
            raise DocumentProcessingError("document_validation_failed") from error
        raise
    if detected_mime != mime_type:
        raise DocumentProcessingError("document_type_mismatch")

    extension = filename.rsplit(".", 1)[-1].lower()
    try:
        if extension == "txt":
            text = data.decode("utf-8", errors="strict")
            pages = [ExtractedPage(number=None, text=text)]
        elif extension == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data), strict=False)
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentProcessingError("pdf_page_limit_exceeded")
            pages = [
                ExtractedPage(number=index + 1, text=page.extract_text() or "")
                for index, page in enumerate(reader.pages)
            ]
        elif extension == "docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            pages = [ExtractedPage(number=None, text="\n".join(lines))]
        else:
            raise DocumentProcessingError("unsupported_document_type")
    except DocumentProcessingError:
        raise
    except Exception as error:
        raise DocumentProcessingError("document_parse_failed") from error

    if sum(len(page.text) for page in pages) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentProcessingError("document_text_limit_exceeded")
    if not any(page.text.strip() for page in pages):
        raise DocumentProcessingError("document_contains_no_extractable_text")
    return pages


def chunk_pages(pages: list[ExtractedPage]) -> list[TextChunk]:
    chunks: list[TextChunk] = []
    for page in pages:
        text = re.sub(r"\s+", " ", page.text).strip()
        start = 0
        while start < len(text):
            end = min(start + MAX_CHUNK_CHARACTERS, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + MAX_CHUNK_CHARACTERS // 2, end)
                if boundary > start:
                    end = boundary
            content = text[start:end].strip()
            if content:
                chunks.append(
                    TextChunk(
                        chunk_index=len(chunks),
                        content=content,
                        page_start=page.number,
                        page_end=page.number,
                    )
                )
            if end >= len(text):
                break
            start = max(start + 1, end - CHUNK_OVERLAP_CHARACTERS)
            if len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
                raise DocumentProcessingError("document_chunk_limit_exceeded")

    if not chunks:
        raise DocumentProcessingError("document_contains_no_extractable_text")
    if len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
        raise DocumentProcessingError("document_chunk_limit_exceeded")
    return chunks


def validate_embeddings(vectors: list[list[float]], expected_count: int) -> int:
    if len(vectors) != expected_count or expected_count == 0:
        raise DocumentProcessingError("embedding_response_count_mismatch")
    dimension = len(vectors[0])
    if not 1 <= dimension <= 16_000:
        raise DocumentProcessingError("embedding_dimension_invalid")
    for vector in vectors:
        if (
            len(vector) != dimension
            or any(not math.isfinite(value) for value in vector)
            or not any(value != 0 for value in vector)
        ):
            raise DocumentProcessingError("embedding_vector_invalid")
    return dimension
